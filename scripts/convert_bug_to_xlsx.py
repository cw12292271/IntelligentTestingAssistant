"""把 bug 记录 txt 转成 xlsx，保留所有字段。"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from pathlib import Path

def parse_bugs(txt_path: str):
    """解析 txt 里的 Bug 记录，提取所有字段。"""
    with open(txt_path, encoding="utf-8") as f:
        content = f.read()

    # 按 "## BUG-" 分割
    blocks = re.split(r"\n## (BUG-\d+)", content)
    bugs = []
    for i in range(1, len(blocks), 2):
        bug_id = blocks[i]
        body = blocks[i + 1].strip()

        # 提取标题（第一行）
        lines = body.split("\n")
        title = lines[0].strip() if lines else ""

        # 提取各字段
        def get_field(name):
            """提取「字段名：内容」的内容，支持多行。"""
            pattern = rf"{name}：\s*(.+?)(?=\n[A-Z\u4e00-\u9fa5]+：|\Z)"
            m = re.search(pattern, body, re.DOTALL)
            return m.group(1).strip() if m else ""

        bugs.append({
            "编号": bug_id,
            "标题": title,
            "模块": get_field("模块"),
            "等级": get_field("等级"),
            "优先级": get_field("优先级"),
            "发现版本": get_field("发现版本"),
            "复现概率": get_field("复现概率"),
            "环境": get_field("环境"),
            "复现步骤": get_field("复现步骤"),
            "预期结果": get_field("预期结果"),
            "实际结果": get_field("实际结果"),
            "原因分析": get_field("原因分析"),
            "修复方案": get_field("修复方案"),
            "回归用例": get_field("回归用例"),
        })
    return bugs


def bugs_to_xlsx(bugs, xlsx_path: str):
    """写入 Excel，带表头样式。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bug记录"

    headers = ["编号", "标题", "模块", "等级", "优先级", "发现版本", "复现概率",
               "环境", "复现步骤", "预期结果", "实际结果", "原因分析", "修复方案", "回归用例"]
    ws.append(headers)

    # 表头样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    for bug in bugs:
        ws.append([bug[h] for h in headers])

    # 列宽
    widths = [10, 30, 12, 8, 8, 12, 10, 20, 50, 30, 30, 40, 40, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    # 单元格自动换行
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    print(f"已生成：{xlsx_path}，共 {len(bugs)} 条")


if __name__ == "__main__":
    bugs = parse_bugs("asset/rag_docs/02_bug记录.txt")
    bugs_to_xlsx(bugs, "asset/rag_docs/02_bug记录.xlsx")
    print(f"样例：{bugs[0]}")