"""把测试用例 txt 转成 xlsx，保留所有字段。"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

def parse_testcases(txt_path: str):
    """解析 txt 里的测试用例，提取所有字段。"""
    with open(txt_path, encoding="utf-8") as f:
        content = f.read()

    # 按 "### TC-" 分割
    blocks = re.split(r"\n### (TC-[\w-]+)", content)
    cases = []
    for i in range(1, len(blocks), 2):
        case_id = blocks[i]
        body = blocks[i + 1].strip()

        # 标题（第一行）
        lines = body.split("\n")
        title = lines[0].strip() if lines else ""

        def get_field(name):
            pattern = rf"\*\*{name}\*\*：\s*(.+?)(?=\n\*\*|\Z)"
            m = re.search(pattern, body, re.DOTALL)
            return m.group(1).strip() if m else ""

        # 操作步骤：提取"**操作步骤**："到下一个**之间的内容
        steps = ""
        m = re.search(r"\*\*操作步骤\*\*：\s*\n(.*?)(?=\n\*\*|\Z)", body, re.DOTALL)
        if m:
            steps = m.group(1).strip()

        cases.append({
            "用例编号": case_id,
            "用例标题": title,
            "优先级": get_field("优先级"),
            "用例类型": get_field("用例类型"),
            "前置条件": get_field("前置条件"),
            "操作步骤": steps,
            "预期结果": get_field("预期结果"),
            "关联需求": get_field("关联需求"),
            "关联Bug": get_field("关联 Bug"),
        })
    return cases


def cases_to_xlsx(cases, xlsx_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    headers = ["用例编号", "用例标题", "优先级", "用例类型", "前置条件",
               "操作步骤", "预期结果", "关联需求", "关联Bug"]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for case in cases:
        ws.append([case[h] for h in headers])

    widths = [14, 30, 8, 10, 20, 50, 30, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    print(f"已生成：{xlsx_path}，共 {len(cases)} 条")


if __name__ == "__main__":
    cases = parse_testcases("asset/rag_docs/05_测试用例文档.txt")
    cases_to_xlsx(cases, "asset/rag_docs/05_测试用例文档.xlsx")
    print(f"样例：{cases[0]}")